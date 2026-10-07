"""Author the Build64 field guide with current offline-rendered UI figures."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / 'docs/field_guide_assets64'
WORK = ROOT / 'docs/field_guide_authoring/revision64'
NAME = 'GIOS 設備助手 現場快速架站手冊'
FONT = 'Microsoft JhengHei'

def font(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + attr), FONT)

def para(text, style=None):
    p = doc.add_paragraph(style=style)
    for i, part in enumerate(re.split(r'\*\*', text)):
        p.add_run(part).bold = bool(i % 2)
    return p

def heading(text):
    doc.add_paragraph(text, style='Heading 1')

def newpage(title):
    doc.add_page_break()
    heading(title)

def pictures(names, caption, width=2.6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    for i, name in enumerate(names):
        if i:
            p.add_run('    ')
        path = name if isinstance(name, Path) else ASSETS / name
        shape = p.add_run().add_picture(str(path), width=Inches(width))
        shape._inline.docPr.set('descr', caption)
    p = para(caption, 'Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

def link(label, url):
    p = doc.add_paragraph()
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), doc.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r')
    rp = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '235D80'); rp.append(color)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rp.append(u)
    r.append(rp)
    t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)
    return p

doc = Document()
for e in list(doc.styles.element.iter(qn('w:pBdr'))):
    e.getparent().remove(e)
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.27), Inches(11.69)
s.top_margin = s.bottom_margin = Inches(.60)
s.left_margin = s.right_margin = Inches(.65)
s.footer_distance = Inches(.25)
font(doc.styles['Normal'], 11.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
font(doc.styles['Title'], 22, True)
doc.styles['Title'].paragraph_format.space_after = Pt(6)
font(doc.styles['Heading 1'], 17, True)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(8)
font(doc.styles['Heading 2'], 13, True)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
font(doc.styles['Caption'], 9)
doc.styles['Caption'].paragraph_format.space_after = Pt(6)
footer = s.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('GIOS 設備助手  |  2026-10-06 更新  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for r in footer.runs: r.font.size = Pt(8)

doc.add_paragraph('GIOS 設備助手\n現場快速架站手冊', style='Title')
para('適用一對一配置與監控　畫面版本 Android 1.0.11 Build 64', 'Caption')
para('**先記住：一個閘道器配一個 PTU。Wi-Fi 必須連線成功且能上網，資料才能上雲。**')
heading('1  從現場配置開始')
para('1. 閘道器、PTU 與路由器通電。備妥站號、**2.4 GHz Wi-Fi 名稱與密碼**；手機開啟藍牙並允許相關權限。')
para('2. 首頁選**現場配置**，確認右上角為**正式站**。按 **⋮ → 連接模式 → 直連模式（一對一）**，關閉設定後按**檢查並開始**。')
para('3. 若提示閘道器目前為星狀模式，選**改成一對一**；若已是一對一，維持即可。')
pictures(['00_home.png', '02_mode.png'], '左：架站選現場配置，日常監控選查看數據。右：確認直連模式下方顯示一對一。', 2.40)
para('**正常架站不要勾「先離線配置，稍後驗證資料」。**離線配置後仍須補做連網與上傳驗證。', None)
para('尚未安裝 APP 請看第 7 頁。圖中站號、網路名稱與數值是操作範例，請以現場設備為準。', 'Caption')

newpage('2  連上設備並設定 Wi Fi')
para('**手機藍牙是設定工具。閘道器要靠現場 Wi-Fi 上網，才能持續上傳資料。**')
para('1. 靠近目標閘道器，按**藍牙連線**。顯示**已連線**後按燈泡，核對閃燈設備，再按**開始開通**。')
para('2. 手機先連現場 Wi-Fi，可先暫停行動數據、開啟網頁確認此 Wi-Fi 能上網。這是初步檢查，還要完成第 3 頁的閘道器檢查。')
para('3. 設定 Wi-Fi 時，按**使用手機目前的 Wi-Fi**帶入名稱，或手動輸入。確認支援 **2.4 GHz**，輸入密碼後按**儲存並繼續**。')
pictures(['04_connected.png', '05_wifi.png'], '左：藍牙連線後先辨識設備。右：核對 Wi-Fi 名稱、密碼，再儲存並繼續。')
para('**只看到 Wi-Fi 名稱或按過儲存，都不代表連線成功。**請繼續等網路檢查結果。已有正常網路時，APP 可能略過輸入頁。')
para('手機系統儲存的 Wi-Fi 密碼不會自動讀出；APP 只會帶入這支手機曾在 APP 記住的密碼。', 'Caption')

newpage('3  確認可以上雲再選站號')
para('**資料路徑：PTU → 閘道器 → 2.4 GHz Wi-Fi → 網際網路 → 雲端**')
para('1. 網路體檢要確認三項：**閘道器已連上 Wi-Fi、資料送到正式站、資料上傳中**。任何一項未通過，先處理，不要當作架站完成。')
para('2. Wi-Fi 已連上卻沒有上傳：先確認路由器有網際網路、資料目的地為正式站；依提示重新檢查。仍失敗請後台協助。')
para('3. 通過後核對工單：站號正確按**使用此站點**；需要換站按**改用其他站號**。新站輸入工單站號，閘道器編號依 APP 分配。')
pictures(['05b_network.png', '06_station.png'], '左：網路體檢三項都要通過。右：站 50 只是範例，請以工單站號為準。')
para('**手機有 4G／5G，不代表閘道器有網路。**閘道器連上的 Wi-Fi 必須能上網；Wi-Fi 圖示亮著也不能取代資料上傳確認。')
para('若剛修改 Wi-Fi，先等「正在等待閘道器恢復資料上傳」結束，不必連續重設。')

newpage('4  確認眼前的 PTU 再開始配置')
para('**一對一模式由閘道器選定一台 PTU。PTU 就是充電樁端的裝置。**')
para('1. 按**辨識此樁**，看現場 PTU 與閘道器閃燈，確認正是要配置的這一樁。')
para('2. 確認正確後，按**是這台，開始配置**，APP 會接著進行配置與資料驗證。')
para('3. 若別樁亮燈、沒有閃燈或無法確認，**不要按確認**。按 **⋯ → 不是這台？**重新處理，必要時請後台協助。')
pictures(['07_identify.png', '08_confirm.png'], '左：先按辨識此樁並看現場燈號。右：確認是目標設備，才按是這台開始配置。')
para('**不要只憑「最近」、訊號最強或已綁定就確認。**設備身分要與現場實物一致。')

newpage('5  驗證上傳並查看充電數據')
para('1. 保持設備通電與網路正常，等 PTU 收到 **3 筆正常資料**、畫面顯示**開通完成**。核對站號與閘道器編號，標示在設備上。')
para('2. 回首頁選**查看數據**，選擇正確站號與閘道器，開啟**最近資料**。完成頁也可直接按**查看最近資料**。')
para('3. 確認顯示**上傳正常**，並等待資料時間持續更新。若資料一直停在舊時間，就還不能確認持續上雲。')
pictures(['09_verify.png', '12_recent.png'], '左：資料驗證由 1/3 累計到 3/3。右：先看上傳狀態與資料時間，再看充電數值。')
para('**車型**：E-Bike＝PRU Type 1；E-Scooter＝Type 2。未知表示缺值或無法識別。')
para('**System Status / Fault**：Normal＝正常；Charging＝充電中；Warning＝警告；Fault＝故障。異常時記下 **Fault Code** 並回報。')
para('**上傳正常不等於充電設備正常。**功率、效率、電池電壓及充電電流會隨狀態變化；畫面數字不是驗收標準值。', 'Caption')

newpage('6  卡住時先看這一頁')
rows = [
    ('遇到的情況', '現場先做什麼'),
    ('找不到或連不上閘道器', '確認通電、藍牙與權限，靠近設備；先讓其他手機斷開連線，再重新搜尋。'),
    ('Wi-Fi 連線失敗', '確認是 2.4 GHz，核對名稱與密碼大小寫；靠近路由器後重新設定。'),
    ('Wi-Fi 已連線但沒有上傳', '檢查該 Wi-Fi 能否上網、上傳目的地是否為正式站；依畫面重新檢查，仍不通就求助。'),
    ('資料一直停住或驗證逾時', '確認設備供電與網路。依畫面等待或重試；回報停在哪一步，勿把舊資料當作成功上傳。'),
    ('辨識時亮的是別樁', '不要確認。選「⋯ → 不是這台？」重新處理；無法確定就求助。'),
    ('藍牙斷線或 APP 中途關閉', '回到現場配置、靠近原設備，依保留進度選「重新連線並繼續」，核對設備後接續。'),
]
table = doc.add_table(rows=0, cols=2)
table.autofit = False
for n, row in enumerate(rows):
    cells = table.add_row().cells
    for i, text in enumerate(row):
        cells[i].text = text
        cells[i].width = Inches(2.05 if i == 0 else 4.90)
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cells[i].paragraphs:
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                r.font.size = Pt(11)
                r.bold = n == 0
        if n == 0 or n % 2 == 0:
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E4EEF4' if n == 0 else 'F4F6F8')
            cells[i]._tc.get_or_add_tcPr().append(shade)
    table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
table.columns[0].width = Inches(2.05); table.columns[1].width = Inches(4.90)
pr = table._tbl.tblPr
borders = OxmlElement('w:tblBorders')
for edge in ('top','left','bottom','right','insideH','insideV'):
    e = OxmlElement('w:'+edge)
    for k,v in [('val','single'),('sz','5'),('color','D9D9D9')]: e.set(qn('w:'+k),v)
    borders.append(e)
pr.append(borders)
margins = OxmlElement('w:tblCellMar')
for edge in ('top','bottom','left','right'):
    e = OxmlElement('w:'+edge); e.set(qn('w:w'),'110'); e.set(qn('w:type'),'dxa'); margins.append(e)
pr.append(margins)
table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
doc.add_paragraph('請後台協助', style='Heading 2')
para('配置流程中點右上方**耳機圖示**。確認「求助已送達後台」，再查看回覆；若未送達，直接電話聯絡。')
para('回報：**站號、閘道器編號、目前步驟、畫面文字、剛做的操作**。不要傳送 Wi-Fi 密碼。')
doc.add_paragraph('收工前確認', style='Heading 2')
para('□ **一對一模式**，站號與 PTU 身分正確。\n□ **Wi-Fi 已連線且能上網**，網路體檢三項通過。\n□ **PTU 驗證 3/3**，最近資料的時間持續更新。\n□ 已標示設備編號，必要時分享安裝報告。')
para('只完成離線配置、只看到藍牙已連線，或只看到閘道器在線，**都不能取代 PTU 資料上雲驗證**。')

newpage('7  安裝與更新 APP')
doc.add_paragraph('Android', style='Heading 2')
link('點此開啟 Android 最新版下載頁', 'https://github.com/jerrywu-voltraware/gateway-commissioning-releases/releases/latest')
para('1. 下載頁展開 **Assets**，點副檔名為 **.apk** 的檔案，下載後開啟安裝。若系統詢問，允許此次安裝來源。')
para('2. 已安裝 APP 可按 **⋮ → 檢查更新 → 立即更新**。若舊版按鈕無法點選，改從上方連結下載 APK **覆蓋安裝，不用先移除**。')
doc.add_paragraph('iPhone', style='Heading 2')
link('點此開啟 iPhone TestFlight 邀請', 'https://testflight.apple.com/join/pv3Xd7PK')
para('1. 在邀請頁先取得 **TestFlight**。\n2. 回到邀請頁，將 GIOS 現場開通加入 TestFlight，再安裝。\n3. 安裝後開啟 **GIOS 設備助手**；若邀請或版本過期，請向管理人員索取新版。')
pictures([WORK/'install0.png', WORK/'install1.png', WORK/'install2.png'], 'iPhone 安裝順序：取得 TestFlight → 加入測試 → 安裝 GIOS 設備助手。', 1.95)
para('不同手機與 iOS 版本的畫面可能略有不同；現場操作重點相同：**一對一、Wi-Fi 有網路、資料確實上雲**。')

doc.core_properties.title = NAME
doc.core_properties.subject = '一對一配置與監控及Wi-Fi上雲確認'
doc.core_properties.author = 'GIOS'
doc.core_properties.keywords = 'GIOS, field guide, one-to-one, Wi-Fi, Build 64'
out = ROOT / (NAME + '.docx')
doc.save(out)
print(out)
