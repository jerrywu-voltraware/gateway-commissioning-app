"""Build the field guide from one content source; screenshots use fixture data."""
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
NAME = 'GIOS 設備助手 現場快速架站手冊'
ASSETS = ROOT / 'docs' / 'field_guide_assets'
FONT = 'Microsoft JhengHei'

PAGES = [
    {
        'title': '01　出發前準備，開啟 APP',
        'intro': '本手冊適用：**一個閘道器配一個 PTU（直連／一對一）**。PTU 是充電樁端的裝置。',
        'steps': [
            '確認閘道器、PTU 與現場 Wi-Fi 路由器已通電；備妥**工單站號與 Wi-Fi 密碼**。',
            '手機開啟藍牙，允許 APP 要求的附近裝置／位置權限；確認右上角是**正式站**。',
            '第一次使用：點 **⋮ → 連接模式 → 直連模式**，回首頁後按**檢查並開始**。',
        ],
        'images': ['01_start.png', '02_mode.png'],
        'caption': '左：按「檢查並開始」。右：一對一安裝選「直連模式」。',
        'tip': '正常連網施工先不要勾「先離線配置，稍後驗證資料」。星狀（一對多）案場請依現場指定流程。',
    },
    {
        'title': '02　連上正確的閘道器',
        'intro': '先確認設備，再開始配置。**點卡片只會選取，連線要按按鈕。**',
        'steps': [
            '靠近本樁的閘道器，在對應卡片按**藍牙連線**，等標籤變成**已連線**。',
            '按卡片上的**燈泡**，看現場哪一台閃燈，確認是目標設備。',
            '確認正確後按**開始開通**。若選錯，先按**斷開**，再連另一台。',
        ],
        'images': ['03_gateway.png', '04_connected.png'],
        'caption': '左：先按「藍牙連線」。右：連線成功後可按燈泡，再「開始開通」。',
        'tip': '「最近」與 dBm 訊號強弱只能幫助找設備，不能當作身分確認。若燈泡提示尚無可辨識 PTU，先核對閘道器標示，後續仍須辨識 PTU。',
    },
    {
        'title': '03　設定現場 Wi-Fi',
        'intro': '閘道器只能連 **2.4 GHz Wi-Fi**。手機先連上現場網路；若 APP 詢問是否重設，按**是，重設 Wi-Fi**，或在網路體檢按**設定 Wi-Fi／重設 Wi-Fi**進入此頁。',
        'steps': [
            '按**使用手機目前的 Wi-Fi**帶入名稱；必要時選**手動輸入其他網路**。',
            '第一次請輸入密碼；按**眼睛**可顯示／隱藏。勾**記住密碼（僅限這支手機）**，成功連線後才記住。',
            '按**儲存並繼續**，等 APP 確認連線與資料上傳，再進入下一步。',
        ],
        'images': ['05_wifi.png'],
        'caption': '圖為「保留原站點、只更新 Wi-Fi」的範例；新機同樣先帶入網路名稱。',
        'tip': '手機系統的 Wi-Fi 密碼不會自動讀出；同一支手機在 APP 記住的密碼才能帶入。若原網路正常，流程可能直接略過此頁。',
    },
    {
        'title': '04　依工單決定站號',
        'intro': '**這裡有兩個正確選項，依實際施工地點選擇。** 圖中的站 50 只是範例。',
        'steps': [
            '目前站號就是工單站號：按**使用此站點**。',
            '搬到別站、站號不符：按**改用其他站號**。新機若直接出現輸入欄，也填工單站號；出現**確定是新站？**時核對後確認。',
            '若剛更新 Wi-Fi，先等「正在等待閘道器恢復資料上傳」結束；不用一直重設。',
        ],
        'images': ['06_station.png'],
        'caption': '沿用原站或改站號都可以；請以工單與現場位置為準。',
        'tip': '站號不確定時，先請後台確認，勿直接照抄範例站號。完成時再核對並標示實際站號、閘道器編號。',
    },
    {
        'title': '05　辨識 PTU，確認是眼前這一樁',
        'intro': '**一定要看現場燈號。** 不要只憑訊號最強或 APP 顯示已綁定就直接確認。',
        'steps': [
            '找到 PTU 後，按**辨識此樁**，看眼前充電樁上的 PTU／閘道器燈號。',
            '確定目標正確後，按**是這台，開始配置**。APP 會接著配置並驗證資料。',
            '若亮的是別樁或無法確認：按底部 **⋯ → 不是這台？**，依畫面重新處理；不確定就求助。',
        ],
        'images': ['07_identify.png', '08_confirm.png'],
        'caption': '左：先「辨識此樁」。右：確認實體設備後，才按「是這台，開始配置」。',
        'tip': '畫面中的 MAC 是裝置識別碼。若需請後台協助，可回報站號、閘道器編號與 MAC 後 4 碼。',
    },
    {
        'title': '06　等驗證完成，再收工',
        'intro': '配置後會自動驗證。**每台 PTU 收到 3 筆正常資料，才算通過。**',
        'steps': [
            '保持設備通電並等進度走完；「剩餘秒數」是等待倒數，不是網路延遲。',
            '看到**開通完成**後，核對站號、閘道器編號與 PTU，並把站號／閘道器編號標在機殼上。',
            '可按**查看最近資料**確認上傳；展開**安裝報告**可分享。最後按**完成**，或**配置下一台**。',
        ],
        'images': ['09_verify.png', '10_done.png'],
        'caption': '左：等待 1/3 → 3/3；未計入資料可展開查看。右：完成後核對並標示設備。',
        'tip': '偶爾一筆「未計入」先等後續正常資料；持續不前進或逾時再求助。「先完成配置／閘道器配置完成」但未驗證 PTU，不代表整樁驗收完成。',
    },
    {
        'title': '07　卡住時，請後台協助',
        'intro': '**先留在目前畫面，不必急著重做。** APP 沒有錯誤也可以主動求助。',
        'steps': [
            '點頂部的**耳機圖示**，開啟**請後台協助**，確認出現**求助已送達後台**。',
            '看後台回覆，依指引操作。若已關閉面板，要再點耳機查看；不要只等背景通知。',
            '排除後按**已解決**；仍卡住按**仍需協助**。未送達時按**重新傳送**，或直接電話聯絡。',
        ],
        'images': ['11_help.png'],
        'caption': '範例：後台回覆站號指引；依現場情況操作，再回報是否解決。',
        'tip': '電話請說：「站號／閘道器編號、MAC 後 4 碼、目前第幾步、畫面文字、剛按了什麼」。求助資訊不會傳送 Wi-Fi 密碼。',
    },
    {
        'title': '08　現場速查',
        'intro': '遇到問題先做右欄動作；仍無法繼續，就使用第 7 頁的後台求助。',
        'table': [
            ['遇到的情況', '先這樣處理'],
            ['找不到／連不上閘道器', '確認通電、手機藍牙與權限，靠近後再試。不要連續切換不同卡片；換設備先斷開。'],
            ['藍牙中途斷線', '先靠近等自動重連；若出現按鈕，按「重新連線並繼續」或「重試重新連線」。'],
            ['Wi-Fi 連不上', '核對 2.4 GHz 網路、名稱與密碼；用眼睛確認輸入。有舊密碼時可改寫或「忘記已存密碼」。'],
            ['一直等待資料／未計入', '先等後續資料；持續未增加或逾時就求助，告知停在哪一步。後台在線不等於 PTU 已驗證。'],
            ['辨識時別樁亮燈', '不要確認配置。選「⋯ → 不是這台？」；無法確定時請後台協助。'],
            ['APP 關閉後要接續', '重新開啟；若顯示保留進度，選「重新連線並繼續」，核對設備後再接續。'],
        ],
        'sections': [
            ('只改 Wi-Fi，不搬站', [
                '連上原閘道器並進入流程；在網路體檢選**重設 Wi-Fi**，或在站點頁選**改用其他 Wi-Fi**。',
                '依第 3 頁帶入名稱、填密碼、儲存；等閘道器恢復上傳。',
                '回到站點頁後，站號正確就選**使用此站點**，並跟著 APP 完成後續確認。',
            ]),
            ('收工前 3 件事', [
                '**身分對**：站號、閘道器編號、PTU 與工單／現場一致。',
                '**資料對**：PTU 已通過 3 筆正常資料驗證；完成頁顯示資料持續上傳。',
                '**標示好**：機殼標示完成，必要時分享安裝報告給後台。',
            ]),
        ],
    },
]


def font_settings(style, size, bold=False, color='000000'):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + attr), FONT)


def rich(p, text):
    for i, part in enumerate(re.split(r'\*\*', text)):
        r = p.add_run(part)
        if i % 2:
            r.bold = True
    return p


def para(doc, text, style=None):
    return rich(doc.add_paragraph(style=style), text)


def add_images(doc, filenames):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    width = 2.62 if len(filenames) > 1 else 2.91
    for index, filename in enumerate(filenames):
        if index:
            p.add_run('     ')
        shape = p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
        shape._inline.docPr.set('descr', 'Build 51 APP 介面範例：' + filename)


def configure_table(table):
    table.autofit = False
    table.columns[0].width = Inches(1.75)
    table.columns[1].width = Inches(5.15)
    pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')]:
            e.set(qn('w:' + key), value)
        borders.append(e)
    pr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge in ('top', 'bottom', 'left', 'right'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:w'), '90' if edge in ('top', 'bottom') else '120')
        e.set(qn('w:type'), 'dxa')
        margins.append(e)
    pr.append(margins)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.width = Inches(1.75 if j == 0 else 5.15)
            if i == 0 or i % 2 == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E8EFF3' if i == 0 else 'F5F7F9')
                cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.08
                for r in p.runs:
                    r.font.size = Pt(10.5)
                    r.bold = i == 0
        cant_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(cant_split)
    repeat = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)


def build_docx():
    doc = Document()
    # The bundled default template may carry a decorative Title border.
    for element in doc.styles.element.iter(qn('w:pBdr')):
        element.getparent().remove(element)
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)
    s.top_margin, s.bottom_margin = Inches(.57), Inches(.55)
    s.left_margin = s.right_margin = Inches(.63)
    s.header_distance, s.footer_distance = Inches(.24), Inches(.24)
    font_settings(doc.styles['Normal'], 11)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.1
    doc.styles['Normal'].paragraph_format.space_after = Pt(5)
    font_settings(doc.styles['Title'], 23, True)
    doc.styles['Title'].paragraph_format.space_after = Pt(4)
    font_settings(doc.styles['Heading 1'], 18, True)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(9)
    font_settings(doc.styles['Heading 2'], 13, True)
    font_settings(doc.styles['Caption'], 9, False, '4A5560')
    doc.styles['Caption'].paragraph_format.space_after = Pt(5)
    header = s.header.paragraphs[0]
    header.text = 'GIOS 設備助手　｜　現場操作'
    for r in header.runs:
        r.font.size = Pt(8)
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('1.0.11 Build 51　•　2026-10-04　｜　')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for r in footer.runs:
        r.font.size = Pt(8)

    for i, page in enumerate(PAGES):
        if i:
            doc.add_page_break()
        else:
            doc.add_paragraph('GIOS 設備助手\n現場快速架站手冊', style='Title')
            para(doc, 'Android 1.0.11 Build 51｜一對一安裝｜2026-10-04', 'Caption')
        doc.add_paragraph(page['title'], style='Heading 1')
        para(doc, page['intro'])
        for n, step in enumerate(page.get('steps', []), 1):
            para(doc, f'{n}. {step}')
        if page.get('images'):
            add_images(doc, page['images'])
            cap = para(doc, page['caption'], 'Caption')
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para(doc, '**提醒｜**' + page['tip'])
        if i == 0:
            para(doc, '圖例：橘框＝操作位置。畫面以本版 APP 介面與範例資料製作，非現場實測紀錄；站號、Wi-Fi 名稱與 MAC 以實際設備為準。', 'Caption')
        if 'table' in page:
            table = doc.add_table(rows=0, cols=2)
            for row in page['table']:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = text
            configure_table(table)
            for title, lines in page['sections']:
                doc.add_paragraph(title, style='Heading 2')
                for n, line in enumerate(lines, 1):
                    para(doc, f'{n}. {line}')
    doc.core_properties.title = NAME
    doc.core_properties.subject = '一對一閘道器與 PTU 現場安裝快速操作'
    doc.core_properties.author = 'GIOS'
    doc.core_properties.keywords = 'GIOS, PTU, gateway, field guide, Build 51'
    path = ROOT / f'{NAME}.docx'
    doc.save(path)
    return path


def build_markdown():
    lines = [f'# {NAME}', '', 'Android 1.0.11 Build 51｜一對一安裝｜2026-10-04', '',
             '> 圖例：橘框＝操作位置。畫面以本版 APP 介面與範例資料製作，非現場實測紀錄；站號、Wi-Fi 名稱與 MAC 以實際設備為準。', '']
    for page in PAGES:
        lines.extend(['## ' + page['title'], '', page['intro'], ''])
        lines.extend(f'{i}. {step}' for i, step in enumerate(page.get('steps', []), 1))
        lines.append('')
        for filename in page.get('images', []):
            lines.extend([f"![{page['caption']}](docs/field_guide_assets/{filename})", ''])
        if 'caption' in page:
            lines.extend([page['caption'], '', '**提醒｜**' + page['tip'], ''])
        if 'table' in page:
            for i, row in enumerate(page['table']):
                lines.append('| ' + ' | '.join(row) + ' |')
                if i == 0:
                    lines.append('| --- | --- |')
            lines.append('')
            for title, items in page['sections']:
                lines.extend(['### ' + title, ''])
                lines.extend(f'{i}. {line}' for i, line in enumerate(items, 1))
                lines.append('')
    path = ROOT / f'{NAME}.md'
    path.write_text('\n'.join(lines), encoding='utf-8')
    return path


if __name__ == '__main__':
    for output in (build_docx(), build_markdown()):
        print(output)
